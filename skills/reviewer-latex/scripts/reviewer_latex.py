#!/usr/bin/env python3
"""Convert reviewer-report drafts to a compact LaTeX report and optional PDF."""

from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from build_support import (
    atomic_copy,
    choose_engine,
    dependency_report,
    enforce_strict,
    font_setup,
    latexmk_command,
    project_version,
    warning_summary,
)
from build_script import write_build_script


SKILL_DIR = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = SKILL_DIR / "assets" / "reviewer_report_template.tex"

SECTION_ALIASES = {
    "overall": {
        "overall",
        "overall assessment",
        "overall comment",
        "overall comments",
        "general comment",
        "general comments",
        "summary",
        "comments to author",
        "comments to authors",
    },
    "major": {
        "major",
        "major concern",
        "major concerns",
        "major comment",
        "major comments",
        "major issue",
        "major issues",
        "major revision concern",
        "major revision concerns",
        "critical concern",
        "critical concerns",
    },
    "minor": {
        "minor",
        "minor concern",
        "minor concerns",
        "minor comment",
        "minor comments",
        "minor issue",
        "minor issues",
    },
    "notes": {
        "note",
        "notes",
        "note to editor",
        "notes to editor",
        "confidential comment",
        "confidential comments",
        "confidential comments to editor",
        "editor note",
        "editor notes",
    },
}

LABEL_ALIASES = {
    "journal": "journal",
    "journal abbreviation": "journal_short",
    "journal short": "journal_short",
    "journal short name": "journal_short",
    "review id": "review_id",
    "reviewer id": "review_id",
    "review number": "review_id",
    "manuscript number": "review_id",
    "manuscript id": "review_id",
    "ms number": "review_id",
    "paper title": "title",
    "manuscript title": "title",
    "manuscript": "title",
    "title": "title",
    "decision": "decision",
    "recommendation": "decision",
}

DEFAULT_TEXT = {
    "overall": r"\textit{No overall comments provided.}",
    "major": r"\textit{No major concerns specified.}",
    "minor": r"\textit{No minor concerns specified.}",
}

JOURNAL_SHORT_NAMES = {
    "npj science of food": "npj Sci Food",
}

SUPERSCRIPT_CHARS = {
    "⁰": "0",
    "¹": "1",
    "²": "2",
    "³": "3",
    "⁴": "4",
    "⁵": "5",
    "⁶": "6",
    "⁷": "7",
    "⁸": "8",
    "⁹": "9",
    "⁺": "+",
    "⁻": "-",
}


@dataclass
class ReviewPoint:
    text: str
    body: list[str] = field(default_factory=list)
    is_heading: bool = False


@dataclass
class Report:
    journal: str = "Journal not specified"
    journal_short: str = ""
    review_id: str = "Not specified"
    title: str = "Untitled manuscript"
    decision: str = "Not specified"
    sections: dict[str, list[str]] = field(
        default_factory=lambda: {"overall": [], "major": [], "minor": [], "notes": []}
    )


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt", ".tex"}:
        return read_plain_text(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix == ".pdf":
        return read_pdf(path)
    raise SystemExit(f"Unsupported input type: {suffix}. Use .md, .txt, .docx, or text-based .pdf.")


def read_plain_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit("DOCX input requires python-docx. Run with the Codex bundled Python runtime.") from exc

    doc = Document(path)
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = normalize_space(paragraph.text)
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [normalize_space(cell.text) for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        pdftotext = shutil.which("pdftotext")
        if not pdftotext:
            raise SystemExit(
                "PDF input requires either the optional pypdf package or the pdftotext command. "
                "Scanned PDFs are not supported."
            )
        result = subprocess.run(
            [pdftotext, str(path), "-"], capture_output=True, text=True, check=False
        )
        if result.returncode != 0:
            raise SystemExit(f"pdftotext failed for {path}: {result.stderr.strip()}")
        text = result.stdout.strip()
    else:
        reader = PdfReader(str(path))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if not text:
        raise SystemExit("No extractable text found in PDF. OCR/scanned PDFs are not supported by this skill.")
    return text


def normalize_space(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def clean_markdown_heading(line: str) -> str:
    line = line.strip()
    line = re.sub(r"^#{1,6}\s*", "", line)
    line = re.sub(r"^\*\*(.+)\*\*$", r"\1", line)
    line = re.sub(r"^__(.+)__$", r"\1", line)
    return line.strip()


def clean_content_line(line: str) -> str:
    line = line.rstrip()
    leading = line[: len(line) - len(line.lstrip())]
    content = clean_markdown_heading(line.lstrip())
    return f"{leading}{content}" if content else ""


def trim_blank_lines(lines: list[str]) -> list[str]:
    trimmed = list(lines)
    while trimmed and not trimmed[0].strip():
        trimmed.pop(0)
    while trimmed and not trimmed[-1].strip():
        trimmed.pop()
    return trimmed


def normalize_key(text: str) -> str:
    text = clean_markdown_heading(text).lower()
    text = re.sub(r"^\s*(section\s+)?[a-z]?\d+(\.\d+)*[.)]?\s+", "", text)
    text = re.sub(r"^\s*[ivxlcdm]+[.)]\s+", "", text)
    text = text.strip(" :-–—\t")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def classify_heading(line: str) -> str | None:
    key = normalize_key(line)
    for section, aliases in SECTION_ALIASES.items():
        if key in aliases:
            return section
    return None


def parse_metadata_line(line: str) -> tuple[str, str] | None:
    raw = clean_markdown_heading(line)
    raw = re.sub(r"^\*\*(.+?):\*\*\s*(.+)$", r"\1: \2", raw)
    match = re.match(r"^([A-Za-z][A-Za-z /_-]{1,35})\s*:\s*(.+)$", raw)
    if not match:
        return None
    label = normalize_key(match.group(1))
    field = LABEL_ALIASES.get(label)
    if not field:
        return None
    value = normalize_space(match.group(2).strip(" *"))
    return field, value


def parse_report(text: str) -> Report:
    report = Report()
    current = "overall"

    for raw_line in text.splitlines():
        stripped = raw_line.strip()
        if not stripped:
            if report.sections[current] and report.sections[current][-1] != "":
                report.sections[current].append("")
            continue

        metadata = parse_metadata_line(stripped)
        if metadata:
            field_name, value = metadata
            if value and getattr(report, field_name) in {
                "",
                "Journal not specified",
                "Not specified",
                "Untitled manuscript",
            }:
                setattr(report, field_name, value)
            continue

        section = classify_heading(stripped)
        if section:
            current = section
            continue

        report.sections[current].append(clean_content_line(raw_line))

    trim_sections(report)
    return report


def trim_sections(report: Report) -> None:
    for key, lines in report.sections.items():
        while lines and not lines[0].strip():
            lines.pop(0)
        while lines and not lines[-1].strip():
            lines.pop()


def apply_overrides(report: Report, args: argparse.Namespace) -> None:
    for attr, value in {
        "journal": args.journal,
        "journal_short": args.journal_short,
        "review_id": args.review_id,
        "title": args.title,
        "decision": args.decision,
    }.items():
        if value:
            setattr(report, attr, value)


def has_cjk(text: str) -> bool:
    return any(
        "\u3400" <= char <= "\u4dbf"
        or "\u4e00" <= char <= "\u9fff"
        or "\uf900" <= char <= "\ufaff"
        for char in text
    )


def escape_latex(text: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(char, char) for char in text)


def latex_url(url: str) -> str:
    safe = url.replace("\\", "/").replace("{", "%7B").replace("}", "%7D")
    return rf"\url{{{safe}}}"


def escape_plain_text(text: str) -> str:
    parts: list[str] = []
    pos = 0
    superscript_chars = "".join(SUPERSCRIPT_CHARS)
    pattern = re.compile(rf"([A-Za-z][A-Za-z0-9]*)([{superscript_chars}]+)")
    for match in pattern.finditer(text):
        parts.append(escape_latex(text[pos : match.start()]))
        base = escape_latex(match.group(1))
        exponent = "".join(SUPERSCRIPT_CHARS[char] for char in match.group(2))
        parts.append(rf"{base}$^{{{exponent}}}$")
        pos = match.end()
    parts.append(escape_latex(text[pos:]))
    return "".join(parts)


def escape_with_urls(text: str) -> str:
    parts: list[str] = []
    pos = 0
    token_pattern = re.compile(r"https?://[^\s<>)]+|\$\$[^$\n]+?\$\$|\$[^$\n]+?\$")
    for match in token_pattern.finditer(text):
        parts.append(escape_plain_text(text[pos : match.start()]))
        token = match.group(0)
        if token.startswith("http"):
            parts.append(latex_url(token))
        else:
            parts.append(token)
        pos = match.end()
    parts.append(escape_plain_text(text[pos:]))
    return "".join(parts)


def inline_latex(text: str) -> str:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+?\*\*|__[^_]+?__|\*[^*]+?\*|_[^_]+?_)" )
    rendered: list[str] = []
    pos = 0
    for match in pattern.finditer(text):
        rendered.append(escape_with_urls(text[pos : match.start()]))
        token = match.group(0)
        if token.startswith("`"):
            rendered.append(r"\texttt{" + escape_latex(token[1:-1]) + "}")
        elif token.startswith(("**", "__")):
            rendered.append(r"\textbf{" + escape_with_urls(token[2:-2]) + "}")
        else:
            rendered.append(r"\emph{" + escape_with_urls(token[1:-1]) + "}")
        pos = match.end()
    rendered.append(escape_with_urls(text[pos:]))
    return "".join(rendered)


def local_heading(line: str) -> str | None:
    stripped = clean_markdown_heading(line)
    if len(stripped) > 120:
        return None
    if stripped.endswith((".", ";", ",")):
        return None
    if re.match(r"^(M|m)?\d+(\.\d+)*[.)]?\s+\S+", stripped):
        return stripped
    if re.match(r"^[A-Z][A-Za-z0-9 /,()_-]{8,}$", stripped) and len(stripped.split()) <= 12:
        return stripped
    return None


def render_lines(lines: list[str], default: str) -> str:
    if not lines:
        return default

    output: list[str] = []
    list_state: str | None = None

    def close_list() -> None:
        nonlocal list_state
        if list_state:
            output.append(rf"\end{{{list_state}}}")
            list_state = None

    for raw in lines:
        line = raw.strip()
        if not line:
            close_list()
            output.append("")
            continue

        bullet = re.match(r"^[-*+]\s+(.+)$", line)
        enum = re.match(r"^\d+[.)]\s+(.+)$", line)
        quote = re.match(r"^>\s+(.+)$", line)
        if bullet or enum:
            env = "itemize" if bullet else "enumerate"
            item_text = (bullet or enum).group(1)
            if list_state != env:
                close_list()
                output.append(rf"\begin{{{env}}}")
                list_state = env
            output.append(r"\item " + inline_latex(item_text))
            continue

        if quote:
            close_list()
            output.append(r"\begin{quote}")
            output.append(inline_latex(quote.group(1)))
            output.append(r"\end{quote}")
            continue

        heading = local_heading(line)
        if heading:
            close_list()
            output.append(r"\subsection*{" + inline_latex(heading) + "}")
            continue

        close_list()
        output.append(inline_latex(line))

    close_list()
    return "\n".join(output).strip() or default


def split_numbered_item_text(text: str) -> tuple[str, list[str], bool]:
    text = text.strip()
    match = re.match(r"^(\*\*|__)(.+?)\1\s*(.*)$", text)
    if not match:
        return text, [], False

    heading = clean_point_heading(match.group(2))
    inline_body = match.group(3).strip()
    return heading, [inline_body] if inline_body else [], True


def clean_point_heading(text: str) -> str:
    return clean_markdown_heading(text).strip().rstrip(" .。")


def parse_review_points(lines: list[str]) -> list[ReviewPoint]:
    points: list[ReviewPoint] = []
    current: ReviewPoint | None = None
    found_numbered = False

    for raw in lines:
        if not raw.strip():
            if current and current.body and current.body[-1] != "":
                current.body.append("")
            continue

        match = re.match(r"^(\d+)[.)]\s+(.+)$", raw)
        if match:
            if current:
                current.body = trim_blank_lines(current.body)
                points.append(current)
            found_numbered = True
            text, body, is_heading = split_numbered_item_text(match.group(2))
            current = ReviewPoint(text=text, body=body, is_heading=is_heading)
            continue

        if current:
            current.body.append(raw)

    if current:
        current.body = trim_blank_lines(current.body)
        points.append(current)

    return points if found_numbered else []


def render_review_points(lines: list[str], default: str, start_number: int) -> tuple[str, int]:
    points = parse_review_points(lines)
    if not points:
        return render_lines(lines, default), start_number

    output: list[str] = []
    number = start_number
    for point in points:
        body = trim_blank_lines(point.body)
        if point.is_heading or body:
            heading = clean_point_heading(point.text)
            output.append(rf"\reviewpointheading{{{number}}}{{{inline_latex(heading)}}}")
            if body:
                output.append(r"\begin{reviewpointbody}")
                output.append(render_lines(body, ""))
                output.append(r"\end{reviewpointbody}")
        else:
            output.append(rf"\reviewplainpoint{{{number}}}{{{inline_latex(point.text)}}}")
        number += 1

    return "\n".join(output), number


def journal_short_name(journal: str) -> str:
    normalized = normalize_key(journal)
    if not normalized or normalized == "journal not specified":
        return "Journal"
    return JOURNAL_SHORT_NAMES.get(normalized, journal)


def render_font_setup(report: Report, engine: str) -> str:
    all_text = "\n".join(
        [report.journal, report.journal_short, report.review_id, report.title, report.decision]
        + [line for lines in report.sections.values() for line in lines]
    )
    return font_setup(engine, has_cjk(all_text))


def render_template(
    report: Report, date: str, report_title: str, body_sections: str, engine: str
) -> str:
    template = TEMPLATE_PATH.read_text(encoding="utf-8")
    journal_short = report.journal_short or journal_short_name(report.journal)

    replacements = {
        "<<FONT_SETUP>>": render_font_setup(report, engine),
        "<<REPORT_TITLE>>": inline_latex(report_title),
        "<<DOCUMENT_MARK>>": "Confidential" if report_title.startswith("Confidential") else "",
        "<<JOURNAL>>": inline_latex(report.journal),
        "<<JOURNAL_SHORT>>": inline_latex(journal_short),
        "<<REVIEW_ID>>": inline_latex(report.review_id),
        "<<TITLE>>": inline_latex(report.title),
        "<<DECISION>>": inline_latex(report.decision),
        "<<DATE>>": inline_latex(date),
        "<<BODY_SECTIONS>>": body_sections,
    }
    for token, value in replacements.items():
        template = template.replace(token, value)
    return template


def render_report(report: Report, date: str, engine: str = "xelatex") -> str:
    major, next_number = render_review_points(report.sections["major"], DEFAULT_TEXT["major"], 1)
    minor, _ = render_review_points(report.sections["minor"], DEFAULT_TEXT["minor"], next_number)
    body_sections = "\n\n".join(
        [
            "\\reportsection{Overall Comments}\n" + render_lines(report.sections["overall"], DEFAULT_TEXT["overall"]),
            "\\reportsection{Major Concerns}\n" + major,
            "\\reportsection{Minor Concerns}\n" + minor,
        ]
    )
    return render_template(report, date, "Peer Review Report", body_sections, engine)


def render_editor_notes(report: Report, date: str, engine: str = "xelatex") -> str | None:
    notes = render_lines(report.sections["notes"], "")
    if not notes:
        return None
    body_sections = "\\reportsection{Confidential Notes to Editor}\n" + notes
    return render_template(report, date, "Confidential Notes to Editor", body_sections, engine)


def compile_tex(tex_path: Path, engine: str, strict: bool = False) -> tuple[Path, str]:
    pdf_path = tex_path.with_suffix(".pdf")
    result = subprocess.run(
        latexmk_command(engine, tex_path.name),
        cwd=tex_path.parent,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )
    if result.returncode != 0:
        log_tail = "\n".join(result.stdout.splitlines()[-80:])
        raise SystemExit(f"LaTeX compilation failed for {tex_path}:\n{log_tail}")
    if not pdf_path.exists():
        raise SystemExit(f"LaTeX command finished, but no PDF was created at {pdf_path}")
    final_log = tex_path.with_suffix(".log").read_text(encoding="utf-8", errors="replace")
    if strict:
        enforce_strict(final_log)
    return pdf_path, final_log


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--version", action="version", version=f"Prism {project_version()}")
    parser.add_argument("input", nargs="?", type=Path, help="Input .md, .txt, .docx, or text-based .pdf file")
    parser.add_argument("--outdir", type=Path, help="Final output directory; defaults beside the input. Compilation uses a temporary directory.")
    parser.add_argument("--out", type=Path, help="Author-facing output path; takes precedence over --outdir. With --no-compile, the suffix becomes .tex.")
    parser.add_argument("--basename", help="Output base name without extension.")
    parser.add_argument("--journal", help="Journal name override.")
    parser.add_argument("--journal-short", help="Short journal name for the page header.")
    parser.add_argument("--review-id", help="Review or manuscript ID override.")
    parser.add_argument("--title", help="Paper title override.")
    parser.add_argument("--decision", help="Decision/recommendation override.")
    parser.add_argument("--date", default=dt.date.today().isoformat(), help="Printed date, default: today.")
    parser.add_argument("--engine", choices=("auto", "xelatex", "lualatex"), default="auto", help="LaTeX engine.")
    parser.add_argument("--strict", action="store_true", help="Fail on missing glyphs, unresolved references, or overfull boxes")
    parser.add_argument("--keep-work", action="store_true", help="Keep the temporary build directory")
    parser.add_argument("--check-deps", action="store_true", help="Check commands, packages, and optional input support")
    parser.add_argument("--no-compile", action="store_true", help="Export persistent author/editor .tex files without requiring a TeX compiler.")
    parser.add_argument("--no-build-script", action="store_true", help="Do not create prism-build.sh after a successful build")
    parser.add_argument("--force-build-script", action="store_true", help="Back up and replace a conflicting prism-build.sh")
    return parser


def default_basename(input_path: Path) -> str:
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "-", input_path.stem).strip("._-")
    return normalized or "review"


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    if args.check_deps:
        ok, rows = dependency_report(
            args.engine, ("fancyhdr.sty", "tabularx.sty")
        )
        rows.append(f"DOCX input: {'python-docx available' if shutil.which('python3') and _module_available('docx') else 'optional python-docx missing'}")
        rows.append(f"PDF input: {'available' if _module_available('pypdf') or shutil.which('pdftotext') else 'optional pypdf/pdftotext missing'}")
        print("\n".join(rows))
        return 0 if ok else 1
    if args.input is None:
        raise SystemExit("INPUT is required unless --check-deps is used.")

    input_path = args.input.expanduser().resolve()
    if not input_path.exists():
        raise SystemExit(f"Input file not found: {input_path}")
    if not input_path.is_file():
        raise SystemExit(f"Input is not a file: {input_path}")

    engine = (
        "xelatex" if args.engine == "auto" else args.engine
    ) if args.no_compile else choose_engine(args.engine)
    basename = args.basename or default_basename(input_path)
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", basename):
        raise SystemExit("--basename must contain only letters, digits, dot, underscore, and hyphen.")
    if args.out:
        destination = args.out.expanduser().resolve()
    else:
        output_dir = args.outdir.expanduser().resolve() if args.outdir else input_path.parent
        output_stem = f"{basename}_review" if input_path.suffix.lower() == ".pdf" else basename
        destination = output_dir / f"{output_stem}.pdf"
    if args.no_compile:
        destination = destination.with_suffix(".tex")
    editor_destination = destination.with_name(
        f"{destination.stem}_editor_notes{destination.suffix}"
    )
    if input_path in (destination, editor_destination):
        raise SystemExit("The output paths must not replace the input file.")

    outdir = Path(tempfile.mkdtemp(prefix="prism-review-latex."))
    try:
        text = read_text(input_path)
        report = parse_report(text)
        apply_overrides(report, args)
        tex_path = outdir / f"{basename}.tex"
        tex_path.write_text(render_report(report, args.date, engine), encoding="utf-8")
        editor_notes_tex = render_editor_notes(report, args.date, engine)
        editor_notes_path = None
        if editor_notes_tex:
            editor_notes_path = outdir / f"{basename}_editor_notes.tex"
            editor_notes_path.write_text(editor_notes_tex, encoding="utf-8")
        if args.no_compile:
            atomic_copy(tex_path, destination)
            print(f"Wrote TeX: {destination}")
            if editor_notes_path:
                atomic_copy(editor_notes_path, editor_destination)
                print(f"Wrote TeX: {editor_destination}")
            return 0
        pdf_path, log = compile_tex(tex_path, engine, args.strict)
        warnings = warning_summary(log)
        if warnings:
            print("LaTeX warnings:\n" + "\n".join(warnings[:20]), file=sys.stderr)
        editor_notes_pdf = None
        if editor_notes_path:
            editor_notes_pdf, editor_log = compile_tex(editor_notes_path, engine, args.strict)
            editor_warnings = warning_summary(editor_log)
            if editor_warnings:
                print("Editor-notes LaTeX warnings:\n" + "\n".join(editor_warnings[:20]), file=sys.stderr)
        atomic_copy(pdf_path, destination)
        print(f"Wrote PDF: {destination}")
        if editor_notes_pdf:
            atomic_copy(editor_notes_pdf, editor_destination)
            print(f"Wrote PDF: {editor_destination}")
        elif editor_destination.exists():
            print("Warning: an older editor-notes PDF exists beside the output; it was not removed.", file=sys.stderr)
        if not args.no_build_script:
            script_options = ["--engine", args.engine, "--date", args.date, "--basename", basename]
            for option, value in (
                ("--journal", args.journal),
                ("--journal-short", args.journal_short),
                ("--review-id", args.review_id),
                ("--title", args.title),
                ("--decision", args.decision),
            ):
                if value is not None:
                    script_options.extend([option, value])
            if args.strict:
                script_options.append("--strict")
            write_build_script(
                workflow="reviewer-latex",
                input_path=input_path,
                output_path=destination,
                project_dir=input_path.parent,
                options=script_options,
                force=args.force_build_script,
            )
        return 0
    finally:
        if args.keep_work:
            print(f"Kept work directory: {outdir}", file=sys.stderr)
        else:
            shutil.rmtree(outdir, ignore_errors=True)


def _module_available(name: str) -> bool:
    try:
        __import__(name)
    except ImportError:
        return False
    return True


if __name__ == "__main__":
    sys.exit(main())
